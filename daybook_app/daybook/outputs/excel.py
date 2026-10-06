"""Excel outputs.

daybook_workbook(): one allocation's Daybook - a Summary sheet headed with that
allocation (Last Month / For The Month / To The Month) and one sheet per sub-allocation
in the same row order as the Daybook view.

table_workbook(): any report view's tables, for the "Export Excel" buttons.
"""
import io
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..labels import head_label


MONEY = "#,##0.00"
HEADER_FILL = PatternFill("solid", start_color="FFD9EAF7")
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _number(text):
    try:
        return float(Decimal(str(text)))
    except (InvalidOperation, ValueError):
        return None


def _style_range(sheet, first_row, last_row, last_col):
    for row in sheet.iter_rows(min_row=first_row, max_row=last_row, max_col=last_col):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _header(sheet, row, values):
    for col, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=col, value=value)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def daybook_workbook(daybook, allocation, period_label, summary_rows):
    """summary_rows: [{"label": "20-16", "last": Decimal, "for": Decimal, "to": Decimal}] (+ no total; added here)."""
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    summary["A1"] = "Daybook Summary - Allocation %s" % allocation
    summary["A1"].font = Font(bold=True, size=13)
    summary.merge_cells("A1:D1")
    summary["A2"] = period_label
    _header(summary, 3, ["Sub-Allocation", "Last Month", "For The Month", "To The Month"])
    row = 4
    totals = [Decimal("0")] * 3
    for item in summary_rows:
        values = [item["last"], item["for"], item["to"]]
        summary.cell(row=row, column=1, value=item["label"])
        for col, value in enumerate(values, start=2):
            summary.cell(row=row, column=col, value=float(value)).number_format = MONEY
        totals = [t + v for t, v in zip(totals, values)]
        row += 1
    summary.cell(row=row, column=1, value="TOTAL").font = Font(bold=True)
    for col, value in enumerate(totals, start=2):
        cell = summary.cell(row=row, column=col, value=float(value))
        cell.number_format = MONEY
        cell.font = Font(bold=True)
    _style_range(summary, 3, row, 4)
    for col, width in zip("ABCD", (18, 20, 20, 20)):
        summary.column_dimensions[col].width = width
    summary.freeze_panes = "A4"

    for table in daybook.tables:
        sheet = book.create_sheet(("Allocation " + table.label)[:31])
        columns = 3 + len(table.heads) + 1
        sheet["A1"] = daybook.table_title(table)
        sheet["A1"].font = Font(bold=True, size=13)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
        _header(sheet, 3, ["Sec. - CO6 / CO7", "BILL DESC / PARTY NAME", "UWID"] + [head_label(h) for h in table.heads] + ["Total"])
        r = 4
        for daybook_row in table.rows:
            first, second = daybook_row.label
            sheet.cell(row=r, column=1, value=first + "\n" + second)
            sheet.cell(row=r, column=2, value=daybook_row.description)
            sheet.cell(row=r, column=3, value=daybook_row.uwid_text)
            for col, head in enumerate(table.heads, start=4):
                shown = daybook_row.cell_display(head)
                if shown is None:
                    sheet.cell(row=r, column=col, value="---").alignment = Alignment(horizontal="center")
                elif len(shown) == 1 and _number(shown[0]) is not None:
                    sheet.cell(row=r, column=col, value=_number(shown[0])).number_format = MONEY
                else:
                    sheet.cell(row=r, column=col, value="\n".join(shown))
            total = _number(daybook_row.total_display)
            sheet.cell(row=r, column=columns, value=total if total is not None else daybook_row.total_display).number_format = MONEY
            r += 1
        sheet.cell(row=r, column=1, value="TOTAL")
        for col, head in enumerate(table.heads, start=4):
            value = _number(table.head_total_display(head))
            sheet.cell(row=r, column=col, value=value).number_format = MONEY
        value = _number(table.total_display)
        sheet.cell(row=r, column=columns, value=value if value is not None else table.total_display).number_format = MONEY
        for cell in sheet[r]:
            cell.font = Font(bold=True)
        _style_range(sheet, 3, r, columns)
        sheet.column_dimensions["A"].width = 26
        sheet.column_dimensions["B"].width = 42
        sheet.column_dimensions["C"].width = 16
        for col in range(4, columns + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 15
        sheet.freeze_panes = "D4"

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def table_workbook(title, tables):
    """tables: [ReportTable]. One sheet per table."""
    book = Workbook()
    book.remove(book.active)
    for index, table in enumerate(tables, start=1):
        sheet = book.create_sheet(("%d %s" % (index, table.title))[:31].replace("/", "-").replace(":", " "))
        sheet["A1"] = title
        sheet["A1"].font = Font(bold=True, size=13)
        sheet["A2"] = table.title
        sheet["A2"].font = Font(bold=True)
        _header(sheet, 4, [c.label for c in table.columns])
        r = 5
        for row in table.rows + ([table.total] if table.total else []):
            for col, (column, value) in enumerate(zip(table.columns, row), start=1):
                if column.money and value is not None and value != "":
                    cell = sheet.cell(row=r, column=col, value=float(value))
                    cell.number_format = MONEY
                else:
                    sheet.cell(row=r, column=col, value=None if value is None else str(value))
            if table.total and row is table.total:
                for cell in sheet[r]:
                    cell.font = Font(bold=True)
            r += 1
        _style_range(sheet, 4, r - 1, len(table.columns))
        for col, column in enumerate(table.columns, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = 18 if column.money else max(14, min(40, len(column.label) + 6))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
